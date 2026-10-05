"""Self-check: hit every endpoint with real generated files. Run: python selfcheck.py"""
import io
import zipfile

import fitz
from fastapi.testclient import TestClient
from PIL import Image

from app import app

c = TestClient(app)
ok = fail = 0


def chk(name, cond, extra=""):
    global ok, fail
    extra = str(extra).encode("ascii", "replace").decode()
    if cond:
        ok += 1
        print(f"  PASS  {name} {extra}")
    else:
        fail += 1
        print(f"  FAIL  {name} {extra}")


def make_pdf(n=3, with_image=True):
    d = fitz.open()
    for i in range(n):
        p = d.new_page(width=595, height=842)
        p.insert_text((72, 100), f"Page {i + 1} - PDFCUY test text", fontsize=14)
    if with_image:
        im = Image.new("RGB", (900, 700))
        px = im.load()
        for y in range(700):
            for x in range(0, 900, 3):
                px[x, y] = ((x + y) % 256, (x * 2) % 256, (y * 3) % 256)
        b = io.BytesIO()
        im.save(b, "PNG")
        d[0].insert_image(fitz.Rect(72, 150, 520, 500), stream=b.getvalue())
    out = d.tobytes()
    d.close()
    return out


def make_img(w=400, h=300, fmt="PNG"):
    b = io.BytesIO()
    Image.new("RGB", (w, h), (90, 140, 240)).save(b, fmt)
    return b.getvalue()


def pages_of(data):
    d = fitz.open(stream=data, filetype="pdf")
    n = d.page_count
    d.close()
    return n


PDF = make_pdf(3)
print(f"\nfixture: {len(PDF)} bytes, 3 halaman\n")

print("index + info")
r = c.get("/")
chk("GET /", r.status_code == 200 and "PDFCUY" in r.text)
r = c.post("/api/info", files={"file": ("a.pdf", PDF, "application/pdf")})
chk("info", r.status_code == 200 and r.json()["pages"] == 3, r.text[:60])

print("merge")
r = c.post(
    "/api/merge",
    files=[
        ("files", ("a.pdf", PDF, "application/pdf")),
        ("files", ("b.pdf", make_pdf(2), "application/pdf")),
    ],
)
chk("merge 3+2=5", r.status_code == 200 and pages_of(r.content) == 5)
r = c.post("/api/merge", files=[("files", ("a.pdf", PDF, "application/pdf"))])
chk("merge rejects single", r.status_code == 400)

print("split")
r = c.post(
    "/api/split",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"mode": "extract", "pages": "1,3"},
)
chk("extract 1,3", r.status_code == 200 and pages_of(r.content) == 2)
r = c.post(
    "/api/split", files={"file": ("a.pdf", PDF, "application/pdf")}, data={"mode": "each"}
)
z = zipfile.ZipFile(io.BytesIO(r.content)) if r.status_code == 200 else None
chk("each -> zip of 3", z and len(z.namelist()) == 3, z.namelist() if z else r.text[:60])
r = c.post(
    "/api/split",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"mode": "ranges", "pages": "1-2 3"},
)
z = zipfile.ZipFile(io.BytesIO(r.content)) if r.status_code == 200 else None
chk("ranges -> 2 parts", z and len(z.namelist()) == 2)
r = c.post(
    "/api/split",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"pages": "99"},
)
chk("out-of-range rejected", r.status_code == 400)
r = c.post(
    "/api/split",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"pages": "abc"},
)
chk("bad spec rejected", r.status_code == 400)

print("organize")
r = c.post(
    "/api/organize",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"pages": "3,2,1"},
)
chk("reorder keeps 3", r.status_code == 200 and pages_of(r.content) == 3)
r = c.post(
    "/api/organize",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"remove": "2"},
)
chk("remove -> 2", r.status_code == 200 and pages_of(r.content) == 2)
r = c.post(
    "/api/organize",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"remove": "1-3"},
)
chk("remove all rejected", r.status_code == 400)

print("rotate")
r = c.post(
    "/api/rotate", files={"file": ("a.pdf", PDF, "application/pdf")}, data={"angle": 90}
)
d = fitz.open(stream=r.content, filetype="pdf")
chk("rotate 90 applied", r.status_code == 200 and d[0].rotation == 90, f"rot={d[0].rotation}")
d.close()
r = c.post(
    "/api/rotate", files={"file": ("a.pdf", PDF, "application/pdf")}, data={"angle": 45}
)
chk("non-90 rejected", r.status_code == 400)

print("compress")
for lv in ("low", "medium", "high"):
    r = c.post(
        "/api/compress",
        files={"file": ("a.pdf", PDF, "application/pdf")},
        data={"level": lv},
    )
    shrink = 100 - len(r.content) * 100 // len(PDF)
    chk(
        f"compress {lv}",
        r.status_code == 200 and len(r.content) <= len(PDF) and pages_of(r.content) == 3,
        f"-{shrink}% ({len(PDF)}->{len(r.content)})",
    )

print("pdf -> images")
r = c.post(
    "/api/pdf-to-images",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"pages": "1", "dpi": 100, "fmt": "png"},
)
chk(
    "single png",
    r.status_code == 200 and r.content[:4] == b"\x89PNG",
    r.headers.get("content-type"),
)
r = c.post(
    "/api/pdf-to-images",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"dpi": 72, "fmt": "jpg"},
)
z = zipfile.ZipFile(io.BytesIO(r.content)) if r.status_code == 200 else None
chk("multi -> zip of 3 jpg", z and len(z.namelist()) == 3 and z.namelist()[0].endswith(".jpg"))
r = c.post(
    "/api/pdf-to-images",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"dpi": 9999, "pages": "1"},
)
chk("out-of-range dpi rejected with reason",
    r.status_code == 400 and "between 36 and 300" in r.text, r.text[:80])
r = c.post(
    "/api/pdf-to-images",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"dpi": 300, "pages": "1"},
)
chk("max allowed dpi works", r.status_code == 200)

print("images -> pdf")
r = c.post(
    "/api/images-to-pdf",
    files=[
        ("files", ("1.png", make_img(), "image/png")),
        ("files", ("2.jpg", make_img(600, 400, "JPEG"), "image/jpeg")),
    ],
    data={"size": "a4", "margin": 20},
)
chk("2 images -> 2 pages", r.status_code == 200 and pages_of(r.content) == 2)
r = c.post(
    "/api/images-to-pdf", files=[("files", ("x.png", b"not-an-image", "image/png"))]
)
chk("junk image rejected", r.status_code == 400)

print("extract text")
r = c.post("/api/extract-text", files={"file": ("a.pdf", PDF, "application/pdf")})
chk("text found", r.status_code == 200 and "Page 2" in r.text, repr(r.text[:40]))
print("watermark")
r = c.post(
    "/api/watermark",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"text": "RESTRICTED", "opacity": 0.2},
)
d = fitz.open(stream=r.content, filetype="pdf")
chk("watermark drawn", r.status_code == 200 and "RESTRICTED" in d[1].get_text())
d.close()
r = c.post(
    "/api/watermark",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"text": "   "},
)
chk("empty text rejected", r.status_code == 400)

print("protect / unlock")
r = c.post(
    "/api/protect",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"new_password": "secret123"},
)
enc = r.content
d = fitz.open(stream=enc, filetype="pdf")
chk("encrypted", r.status_code == 200 and d.needs_pass)
d.close()
r = c.post("/api/protect", files={"file": ("a.pdf", PDF, "application/pdf")}, data={"new_password": "ab"})
chk("short pw rejected", r.status_code == 400)

r = c.post("/api/info", files={"file": ("e.pdf", enc, "application/pdf")})
chk("locked rejects no-pw", r.status_code == 401)
r = c.post(
    "/api/info", files={"file": ("e.pdf", enc, "application/pdf")}, data={"password": "secret123"}
)
chk("locked opens with pw", r.status_code == 200 and r.json()["pages"] == 3)
r = c.post(
    "/api/info", files={"file": ("e.pdf", enc, "application/pdf")}, data={"password": "wrong"}
)
chk("wrong pw rejected", r.status_code == 401)

r = c.post(
    "/api/unlock", files={"file": ("e.pdf", enc, "application/pdf")}, data={"password": "secret123"}
)
d = fitz.open(stream=r.content, filetype="pdf")
chk("unlocked", r.status_code == 200 and not d.needs_pass and d.page_count == 3)
d.close()
r = c.post("/api/unlock", files={"file": ("a.pdf", PDF, "application/pdf")})
chk("unlock plain rejected", r.status_code == 400)

print("pdf -> docx")
r = c.post(
    "/api/pdf-to-docx",
    files={"file": ("a.pdf", PDF, "application/pdf")},
    data={"pages": "1-2"},
)
chk(
    "docx produced",
    r.status_code == 200 and r.content[:2] == b"PK" and len(r.content) > 2000,
    f"{len(r.content)}B" if r.status_code == 200 else r.text[:70],
)

# ---------- Office ----------
def make_docx():
    import docx

    d = docx.Document()
    d.add_heading("PDFCUY Test Report", 0)
    d.add_paragraph("First paragraph of ordinary body text for testing.")
    d.add_heading("Section Two", level=2)
    t = d.add_table(rows=3, cols=3)
    for i in range(3):
        for j in range(3):
            t.cell(i, j).text = f"sel{i}{j}"
    d.add_paragraph("Closing line of the document.")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def make_xlsx():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Product", "Qty", "Price"])
    for i in range(1, 12):
        ws.append([f"Item {i}", i * 3, i * 1500])
    wb.create_sheet("Notes").append(["note", "second test"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def make_pptx():
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Test Slide Title"
    s.placeholders[1].text = "Point one\nPoint two"
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Second Slide"
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


print("engines")
r = c.get("/api/engines")
has_lo = r.json().get("libreoffice")
chk("engines reported", r.status_code == 200, f"libreoffice={has_lo}")

DOCX, XLSX, PPTX = make_docx(), make_xlsx(), make_pptx()
CSV = b"name,city,age\nAri,Jakarta,29\nBudi,Bandung,34\n"
TXT = ("Ordinary line of text.\n" * 30).encode()

print("office -> pdf")
for nm, blob in [("a.docx", DOCX), ("a.xlsx", XLSX), ("a.pptx", PPTX),
                 ("a.csv", CSV), ("a.txt", TXT)]:
    r = c.post("/api/office-to-pdf", files={"files": (nm, blob, "application/octet-stream")})
    good = r.status_code == 200 and r.content[:4] == b"%PDF" and pages_of(r.content) >= 1
    chk(nm + " -> pdf", good,
        f"{pages_of(r.content)}p via {r.headers.get('X-Engine')}" if r.status_code == 200 else r.text[:70])

r = c.post(
    "/api/office-to-pdf",
    files=[("files", ("a.docx", DOCX, "application/octet-stream")),
           ("files", ("b.xlsx", XLSX, "application/octet-stream"))],
    data={"combine": "true"},
)
chk("multi combine -> 1 pdf", r.status_code == 200 and r.content[:4] == b"%PDF" and pages_of(r.content) >= 2,
    f"{pages_of(r.content)}p" if r.status_code == 200 else r.text[:70])

r = c.post(
    "/api/office-to-pdf",
    files=[("files", ("a.docx", DOCX, "application/octet-stream")),
           ("files", ("b.csv", CSV, "text/csv"))],
)
z = zipfile.ZipFile(io.BytesIO(r.content)) if r.status_code == 200 else None
chk("multi separate -> zip", z and len(z.namelist()) == 2, z.namelist() if z else r.text[:70])

r = c.post("/api/office-to-pdf", files={"files": ("p.pdf", PDF, "application/pdf")})
chk("pdf passthrough", r.status_code == 200 and pages_of(r.content) == 3)

r = c.post("/api/office-to-pdf", files={"files": ("a.xyz", b"junk data here", "application/octet-stream")})
chk("unknown ext rejected", r.status_code == 422, r.text[:60])

r = c.post("/api/office-to-pdf", files={"files": ("a.docx", b"not really a docx", "application/octet-stream")})
chk("corrupt docx rejected", r.status_code == 422, r.text[:60])

if not has_lo:
    r = c.post("/api/office-to-pdf", files={"files": ("old.doc", b"\xd0\xcf\x11\xe0junk", "application/msword")})
    chk("legacy .doc explains need", r.status_code == 422 and "LibreOffice" in r.text, r.text[:60])

print("pdf -> office")
TBL = fitz.open()
_p = TBL.new_page()
for i, row in enumerate([["Name", "City", "Age"], ["Ari", "Jakarta", "29"], ["Budi", "Bandung", "34"]]):
    for j, cell in enumerate(row):
        _p.insert_text((72 + j * 120, 100 + i * 24), cell, fontsize=11)
TBLPDF = TBL.tobytes()
TBL.close()

r = c.post("/api/pdf-to-xlsx", files={"file": ("t.pdf", TBLPDF, "application/pdf")})
if r.status_code == 200:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    flat = " ".join(
        str(cv) for ws in wb.worksheets for rw in ws.iter_rows(values_only=True) for cv in rw if cv
    )
    chk("pdf -> xlsx has data", "Jakarta" in flat, wb.sheetnames)
else:
    chk("pdf -> xlsx has data", False, r.text[:80])

_EMPTY = fitz.open()
_EMPTY.new_page()
r = c.post("/api/pdf-to-xlsx", files={"file": ("blank.pdf", _EMPTY.tobytes(), "application/pdf")})
_EMPTY.close()
chk("empty/scanned pdf -> 422", r.status_code == 422, r.text[:60])

r = c.post("/api/pdf-to-pptx", files={"file": ("a.pdf", PDF, "application/pdf")}, data={"pages": "1-2"})
if r.status_code == 200:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(r.content))
    chk("pdf -> pptx 2 slides", len(prs.slides._sldIdLst) == 2, f"{len(r.content)}B")
else:
    chk("pdf -> pptx 2 slides", False, r.text[:80])

print("office text")
for nm, blob, needle in [("a.docx", DOCX, "PDFCUY Test Report"), ("a.xlsx", XLSX, "Sales"),
                         ("a.pptx", PPTX, "Test Slide Title")]:
    r = c.post("/api/office-text", files={"file": (nm, blob, "application/octet-stream")})
    chk(f"text from {nm}", r.status_code == 200 and needle in r.text, repr(r.text[:45]))

print("convert endpoint")
r = c.post("/api/convert", files={"file": ("a.docx", DOCX, "application/octet-stream")}, data={"to": "pdf"})
chk("convert to pdf", r.status_code == 200 and r.content[:4] == b"%PDF")
r = c.post("/api/convert", files={"file": ("a.docx", DOCX, "application/octet-stream")}, data={"to": "docx"})
chk("same-format rejected", r.status_code == 400)
r = c.post("/api/convert", files={"file": ("a.docx", DOCX, "application/octet-stream")}, data={"to": "../etc"})
chk("path-traversal target rejected", r.status_code == 400, r.text[:60])
r = c.post("/api/convert", files={"file": ("a.docx", DOCX, "application/octet-stream")}, data={"to": "odt"})
chk("odt without LO -> 501" if not has_lo else "odt with LO -> 200",
    (r.status_code == 501) if not has_lo else (r.status_code == 200), r.text[:50] if r.status_code >= 400 else "")

print("guards")
r = c.post("/api/info", files={"file": ("x.pdf", b"i am not a pdf at all", "application/pdf")})
chk("non-pdf rejected", r.status_code == 400)
r = c.post("/api/info", files={"file": ("x.pdf", b"", "application/pdf")})
chk("empty rejected", r.status_code in (400, 422))
r = c.post("/api/info", files={"file": ("x.pdf", b"%PDF-1.7\ntruncated junk", "application/pdf")})
chk("corrupt rejected", r.status_code == 400)

print("hardening: strict enums (no silent fallback)")
PDF_F = {"file": ("a.pdf", PDF, "application/pdf")}
for path, data, needle in [
    ("/api/compress", {"level": "ludicrous"}, "one of"),
    ("/api/split", {"mode": "nonsense"}, "one of"),
    ("/api/pdf-to-images", {"dpi": -5}, "between"),
    ("/api/pdf-to-pptx", {"dpi": 9999}, "between"),
    ("/api/watermark", {"text": "X", "opacity": 99}, "between"),
    ("/api/watermark", {"text": "X", "fontsize": 0}, "between"),
    ("/api/watermark", {"text": "X", "fontsize": -20}, "between"),
]:
    r = c.post(path, files=PDF_F, data=data)
    chk(f"{path} {list(data)[-1]}={list(data.values())[-1]} -> 400",
        r.status_code == 400 and needle in r.text, f"{r.status_code} {r.text[:60]}")
r = c.post("/api/office-to-pdf",
           files=[("files", ("a.pdf", PDF, "application/pdf"))], data={"combine": "maybe"})
chk("combine=maybe rejected", r.status_code == 400 and "one of" in r.text, r.text[:60])

print("hardening: hostile filenames cannot reach the header")
HOSTILE = [
    "../../../../etc/passwd.pdf",
    "..\\..\\windows\\system32\\evil.pdf",
    'a"b;filename="other.pdf',
    "<script>alert(1)</script>.pdf",
    "A" * 300 + ".pdf",
]
for name in HOSTILE:
    r = c.post("/api/compress", files={"file": (name, PDF, "application/pdf")})
    cd = r.headers.get("content-disposition", "")
    plain = cd.split('"')[1] if cd.count('"') >= 2 else ""
    clean = (r.status_code == 200 and plain.isascii() and cd.count('"') == 2
             and not any(x in plain for x in ("..", "/", "\\", "<", ">", ";", "\r", "\n")))
    chk(f"filename {name[:26]!r} neutralised", clean, plain[:46])

r = c.post("/api/compress",
           files={"file": ("\u4e2d\u6587\u6587\u6863.pdf", PDF, "application/pdf")})
cd = r.headers.get("content-disposition", "")
chk("unicode name survives via filename*",
    "filename*=UTF-8''" in cd and "%E4%B8%AD" in cd, cd[:80])
chk("ascii fallback stays ascii", cd.split('"')[1].isascii(), cd.split('"')[1][:40])

print("hardening: security headers on every response kind")
WANT = ("content-security-policy", "x-content-type-options", "x-frame-options",
        "referrer-policy", "permissions-policy", "cross-origin-opener-policy")
for label, resp in [("html", c.get("/")),
                    ("static css", c.get("/static/app.css")),
                    ("download", c.post("/api/compress", files=PDF_F)),
                    ("error", c.post("/api/compress", files=PDF_F, data={"level": "bad"}))]:
    missing = [h for h in WANT if h not in resp.headers]
    chk(f"{label} carries all 6 headers", not missing, missing or "complete")
chk("csp forbids inline script", "'unsafe-inline'" not in c.get("/").headers["content-security-policy"])
chk("no cookie is ever set", "set-cookie" not in c.get("/").headers)

print("hardening: errors stay opaque")
for path, data in [("/api/watermark", {"text": "X", "fontsize": 0}),
                   ("/api/compress", {"level": "nope"}),
                   ("/api/rotate", {"angle": 45})]:
    r = c.post(path, files=PDF_F, data=data)
    leak = any(s in r.text for s in ("Traceback", 'File "', "site-packages", "app.py"))
    chk(f"{path} hides internals", r.status_code == 400 and not leak, r.text[:50])

print(f"\n{ok} passed, {fail} failed")
assert fail == 0, f"{fail} check(s) failed"
print("ALL GREEN")
