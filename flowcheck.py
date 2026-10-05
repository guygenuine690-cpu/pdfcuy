"""Tool-first flow test: pick tool from grid -> drop file -> run -> download.
Needs server running + playwright chromium."""
import os
import tempfile

import docx
import pymupdf
from playwright.sync_api import sync_playwright

ok = fail = 0


def chk(name, cond, extra=""):
    global ok, fail
    ok, fail = (ok + 1, fail) if cond else (ok, fail + 1)
    print(f"  {'PASS' if cond else 'FAIL'}  {name} {extra}")


TMP = tempfile.mkdtemp(prefix="flow_")


def pdf_path(n=3, name="report.pdf"):
    d = pymupdf.open()
    for i in range(n):
        d.new_page(width=595, height=842).insert_text(
            (72, 100), f"Page {i + 1} of the flow test", fontsize=14)
    p = os.path.join(TMP, name)
    d.save(p)
    d.close()
    return p


def docx_path():
    d = docx.Document()
    d.add_heading("Quarterly Report", 0)
    d.add_paragraph("A paragraph for the end to end flow test.")
    p = os.path.join(TMP, "memo.docx")
    d.save(p)
    return p


def png_path(name="shot.png"):
    from PIL import Image
    p = os.path.join(TMP, name)
    Image.new("RGB", (640, 480), (70, 110, 200)).save(p)
    return p


PDF, PDF2, DOCX, PNG1, PNG2 = (pdf_path(), pdf_path(2, "appendix.pdf"),
                               docx_path(), png_path("a.png"), png_path("b.png"))

with sync_playwright() as p:
    b = p.chromium.launch()
    pg = b.new_page(viewport={"width": 1360, "height": 1000}, device_scale_factor=2)
    errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.on("console", lambda m: m.type == "error" and errs.append(m.text))
    pg.goto("http://127.0.0.1:8000", wait_until="networkidle")
    pg.wait_for_selector(".tool", timeout=10000)
    pg.wait_for_timeout(700)

    print("library - tool grid first")
    chk("library visible, workspace hidden",
        pg.is_visible("#library") and pg.is_hidden("#workspace"))
    chk("27 tool cards", pg.locator(".tool").count() == 27, pg.locator(".tool").count())
    chk("4 categories", pg.locator(".sec").count() == 4, pg.locator(".sec").count())
    chk("category nav built", pg.locator("#tnav a").count() == 4)
    chk("format logos on cards", pg.locator(".tool .flow .lg").count() >= 20,
        pg.locator(".tool .flow .lg").count())
    chk("popular badges", pg.locator(".badge").count() == 5, pg.locator(".badge").count())
    chk("engine pill filled", "engine" in (pg.inner_text("#engine") or "").lower()
        or "memory" in pg.inner_text("#engine").lower(), pg.inner_text("#engine"))

    dup = pg.evaluate("""() => {
      const ids = [...document.querySelectorAll('[id]')].map(e => e.id);
      const seen = new Set(), dupes = new Set();
      for (const i of ids) { if (seen.has(i)) dupes.add(i); seen.add(i); }
      return [...dupes];
    }""")
    chk("no duplicate DOM ids", not dup, dup[:5] or "unique")

    # generated svg gradient ids must never shadow an app control id
    leak = pg.evaluate("""() => [...document.querySelectorAll('linearGradient[id]')]
      .map(e => e.id).filter(i => !i.startsWith('fmt'))""")
    chk("gradient ids namespaced", not leak, leak[:5] or "all fmt-prefixed")
    pg.screenshot(path=os.path.join(TMP, "shot-1-library.png"), full_page=True)

    print("search")
    pg.fill("#q", "password")
    pg.wait_for_timeout(400)
    vis = pg.locator(".tool:visible").count()
    chk("search narrows cards", 1 <= vis <= 8, f"{vis} shown")
    pg.fill("#q", "zzzqqq")
    pg.wait_for_timeout(400)
    chk("empty state appears", pg.is_visible(".empty"))
    pg.fill("#q", "")
    pg.wait_for_timeout(400)
    chk("clearing restores all", pg.locator(".tool:visible").count() == 27)

    print("workspace - compress")
    pg.click('.tool[data-k="compress"]')
    pg.wait_for_selector("#workspace:not([hidden])", timeout=5000)
    pg.wait_for_timeout(500)
    chk("workspace shown", pg.is_visible("#workspace") and pg.is_hidden("#library"))
    chk("title is the tool", pg.inner_text("#wtitle") == "Compress PDF", pg.inner_text("#wtitle"))
    chk("url carries tool", pg.url.endswith("#compress"), pg.url.split("/")[-1])
    chk("breadcrumb category", pg.inner_text("#crumbcat") == "Optimize", pg.inner_text("#crumbcat"))
    chk("dropzone waiting", pg.is_visible("#drop"))
    chk("run disabled with no file", pg.is_disabled("#go"))
    chk("run button asks for file", "Choose a file" in pg.inner_text("#gotxt"), pg.inner_text("#gotxt"))
    chk("segmented control rendered", pg.locator(".seg button").count() == 3,
        pg.locator(".seg button").count())
    chk("privacy panel present", pg.locator(".priv li").count() == 3)
    pg.screenshot(path=os.path.join(TMP, "shot-2-workspace.png"), full_page=True)

    print("wrong file type rejected")
    pg.set_input_files("#picker", PNG1)
    pg.wait_for_timeout(500)
    chk("image refused for Compress PDF", pg.is_visible(".note.e"),
        pg.inner_text(".note.e")[:70] if pg.is_visible(".note.e") else "no error")

    print("correct file accepted")
    pg.set_input_files("#picker", PDF)
    pg.wait_for_selector(".flist li", timeout=5000)
    pg.wait_for_timeout(400)
    chk("file row shown", pg.locator(".flist li").count() == 1)
    chk("format logo in row", pg.locator(".flist .lg svg").count() == 1)
    chk("run now enabled", not pg.is_disabled("#go"))
    chk("run labelled by tool", pg.inner_text("#gotxt") == "Compress PDF", pg.inner_text("#gotxt"))
    chk("summary shows size", "KB" in pg.inner_text("#sum") or "B" in pg.inner_text("#sum"))
    chk("advanced hidden by default", pg.is_visible("#advbtn"))
    pg.screenshot(path=os.path.join(TMP, "shot-3-ready.png"), full_page=True)

    print("segmented control works")
    pg.click('.seg button[data-v="high"]')
    pg.wait_for_timeout(250)
    chk("choice registers",
        pg.get_attribute('.seg button[data-v="high"]', "aria-pressed") == "true"
        and pg.input_value("#o_level") == "high", pg.input_value("#o_level"))

    print("run - real download")
    with pg.expect_download(timeout=60000) as dl:
        pg.click("#go")
    out = os.path.join(TMP, "out.pdf")
    dl.value.save_as(out)
    chk("got a pdf back", open(out, "rb").read(4) == b"%PDF",
        f"{os.path.getsize(out)} B as {dl.value.suggested_filename}")
    pg.wait_for_selector(".note.s", timeout=15000)
    chk("success note shown", "Saved" in pg.inner_text(".note.s"), pg.inner_text(".note.s")[:50])

    print("navigation")
    pg.click("#back")
    pg.wait_for_timeout(400)
    chk("back to library", pg.is_visible("#library"))
    pg.click('.tool[data-k="merge"]')
    pg.wait_for_timeout(400)
    pg.keyboard.press("Escape")
    pg.wait_for_timeout(400)
    chk("escape leaves workspace", pg.is_visible("#library"))

    print("merge needs two files")
    pg.click('.tool[data-k="merge"]')
    pg.wait_for_timeout(400)
    pg.set_input_files("#picker", PDF)
    pg.wait_for_timeout(450)
    chk("one file is not enough", pg.is_disabled("#go")
        and "Add 1 more" in pg.inner_text("#gotxt"), pg.inner_text("#gotxt"))
    chk("add-more button offered", pg.is_visible("#addmore"))
    pg.set_input_files("#picker", PDF2)
    pg.wait_for_timeout(450)
    chk("two files enable run", not pg.is_disabled("#go") and pg.locator(".flist li").count() == 2)
    chk("order numbers shown", pg.locator(".flist .ord").count() == 2)
    first = pg.locator(".flist .fn").first.inner_text()
    pg.click('.mini[data-mv="1"][data-d="-1"]')
    pg.wait_for_timeout(300)
    chk("reorder works", pg.locator(".flist .fn").first.inner_text() != first,
        pg.locator(".flist .fn").first.inner_text())
    with pg.expect_download(timeout=60000) as dl2:
        pg.click("#go")
    merged = os.path.join(TMP, "merged.pdf")
    dl2.value.save_as(merged)
    doc = pymupdf.open(merged)
    chk("merged 3+2 pages", doc.page_count == 5, f"{doc.page_count}p")
    doc.close()

    print("office tool filters input")
    pg.click("#back")
    pg.wait_for_timeout(350)
    pg.click('.tool[data-k="office-to-pdf"]')
    pg.wait_for_timeout(400)
    acc = pg.get_attribute("#picker", "accept")
    chk("accept excludes pdf", ".docx" in acc and ".pdf" not in acc, acc[:46])
    pg.set_input_files("#picker", DOCX)
    pg.wait_for_timeout(450)
    chk("docx accepted", pg.locator(".flist li").count() == 1)
    with pg.expect_download(timeout=90000) as dl3:
        pg.click("#go")
    conv = os.path.join(TMP, "conv.pdf")
    dl3.value.save_as(conv)
    chk("docx became pdf", open(conv, "rb").read(4) == b"%PDF", f"{os.path.getsize(conv)} B")

    print("images to pdf")
    pg.click("#back")
    pg.wait_for_timeout(350)
    pg.click('.tool[data-k="images-to-pdf"]')
    pg.wait_for_timeout(400)
    pg.set_input_files("#picker", [PNG1, PNG2])
    pg.wait_for_timeout(450)
    chk("two images queued", pg.locator(".flist li").count() == 2)
    with pg.expect_download(timeout=60000) as dl4:
        pg.click("#go")
    imgs = os.path.join(TMP, "imgs.pdf")
    dl4.value.save_as(imgs)
    doc = pymupdf.open(imgs)
    chk("one page per image", doc.page_count == 2, f"{doc.page_count}p")
    doc.close()

    print("deep link")
    pg.goto("http://127.0.0.1:8000/#protect", wait_until="networkidle")
    pg.wait_for_timeout(700)
    chk("hash opens the tool", pg.is_visible("#workspace")
        and pg.inner_text("#wtitle") == "Protect PDF", pg.inner_text("#wtitle"))
    chk("required field marked", pg.locator(".rq").count() >= 1)

    print("required validation")
    pg.set_input_files("#picker", PDF)
    pg.wait_for_timeout(400)
    pg.fill("#o_new_password", "")
    pg.click("#go")
    pg.wait_for_timeout(500)
    chk("blocks empty required field", pg.is_visible(".note.e"),
        pg.inner_text(".note.e")[:50] if pg.is_visible(".note.e") else "none")

    print("responsive")
    for w, h in ((1360, 1000), (1024, 900), (820, 900), (560, 900), (390, 860)):
        pg.set_viewport_size({"width": w, "height": h})
        pg.wait_for_timeout(300)
        o = pg.evaluate("() => [document.documentElement.scrollWidth, document.documentElement.clientWidth]")
        chk(f"no overflow at {w}px", o[0] <= o[1] + 1, o)

    pg.goto("http://127.0.0.1:8000", wait_until="networkidle")
    pg.set_viewport_size({"width": 390, "height": 860})
    pg.wait_for_timeout(600)
    pg.screenshot(path=os.path.join(TMP, "shot-4-mobile.png"), full_page=True)

    print("dark theme")
    pg.set_viewport_size({"width": 1360, "height": 1000})
    pg.click("#theme")
    pg.wait_for_timeout(450)
    chk("theme toggles", pg.get_attribute("html", "data-theme") == "dark")
    pg.screenshot(path=os.path.join(TMP, "shot-5-dark.png"), full_page=True)

    chk("zero console errors", not errs, errs[:2] or "clean")
    csp = [e for e in errs if "Content Security Policy" in e]
    chk("zero CSP violations", not csp, f"{len(csp)} blocked" if csp else "none")
    b.close()

print(f"\n{ok} passed, {fail} failed")
assert fail == 0, f"{fail} flow check(s) failed"
print("FLOW OK")
