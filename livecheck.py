"""Live HTTP smoke test against a running server. Run: python livecheck.py"""
import io
import json
import urllib.request as u

import docx

B = "http://127.0.0.1:8000"


def post(path, filename, blob, field="files", extra=None):
    bd, parts = "----X7", []
    for k, v in (extra or {}).items():
        parts.append(
            f'--{bd}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode()
        )
    parts.append(
        f'--{bd}\r\nContent-Disposition: form-data; name="{field}"; filename="{filename}"\r\n'
        f"Content-Type: application/octet-stream\r\n\r\n".encode() + blob + b"\r\n"
    )
    parts.append(f"--{bd}--\r\n".encode())
    body = b"".join(parts)
    rq = u.Request(B + path, data=body,
                   headers={"Content-Type": f"multipart/form-data; boundary={bd}"})
    r = u.urlopen(rq, timeout=120)
    return r, r.read()


print("engines:", json.loads(u.urlopen(B + "/api/engines").read()))
h = u.urlopen(B)
page = h.read()
print(f"index: HTTP {h.status}, {len(page)} bytes, title-ok={b'PDFCUY' in page}")

d = docx.Document()
d.add_heading("Live Test", 0)
d.add_paragraph("halo dari uji live lewat HTTP nyata")
buf = io.BytesIO()
d.save(buf)

r, out = post("/api/office-to-pdf", "live.docx", buf.getvalue())
print(f"docx->pdf: HTTP {r.status} {len(out)}B magic={out[:4]!r} "
      f"engine={r.headers.get('X-Engine')} disp={r.headers.get('Content-Disposition')}")
assert out[:4] == b"%PDF"

r, out = post("/api/pdf-to-pptx", "live.pdf", out, field="file")
print(f"pdf->pptx: HTTP {r.status} {len(out)}B magic={out[:2]!r}")
assert out[:2] == b"PK"

print("LIVE OK")
